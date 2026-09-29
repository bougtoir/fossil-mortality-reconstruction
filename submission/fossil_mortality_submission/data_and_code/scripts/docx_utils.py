"""Helpers for docx generation: OMML equations, figure/table insertion."""
from docx.oxml.ns import nsmap, qn
from docx.oxml import parse_xml
from docx.shared import Inches, Pt
import re

M = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'


def omml_eq(runs):
    """runs: list of str OR ('sup', base, exp) OR ('frac', num, den) tuples."""
    parts = []
    for r in runs:
        if isinstance(r, str):
            parts.append(f'<m:r xmlns:m="{M}"><m:t xml:space="preserve">{r}</m:t></m:r>')
        elif r[0] == 'sup':
            parts.append(
                f'<m:sSup xmlns:m="{M}"><m:e><m:r><m:t>{r[1]}</m:t></m:r></m:e>'
                f'<m:sup><m:r><m:t>{r[2]}</m:t></m:r></m:sup></m:sSup>')
        elif r[0] == 'frac':
            parts.append(
                f'<m:f xmlns:m="{M}"><m:num><m:r><m:t>{r[1]}</m:t></m:r></m:num>'
                f'<m:den><m:r><m:t>{r[2]}</m:t></m:r></m:den></m:f>')
        elif r[0] == 'int':
            parts.append(
                f'<m:nary xmlns:m="{M}"><m:naryPr><m:chr m:val="∫"/>'
                f'<m:limLoc m:val="undOvr"/></m:naryPr>'
                f'<m:sub><m:r><m:t>{r[1]}</m:t></m:r></m:sub>'
                f'<m:sup><m:r><m:t>{r[2]}</m:t></m:r></m:sup>'
                f'<m:e><m:r><m:t>{r[3]}</m:t></m:r></m:e></m:nary>')
    xml = (f'<m:oMathPara xmlns:m="{M}" xmlns:w="{W}">'
           f'<m:oMathParaPr><m:jc m:val="center"/></m:oMathParaPr>'
           f'<m:oMath>{"".join(parts)}</m:oMath></m:oMathPara>')
    return parse_xml(xml)


def add_equation(doc, runs):
    p = doc.add_paragraph()
    p._p.append(omml_eq(runs))
    return p


def add_figure(doc, path, caption, width=6.2):
    p = doc.add_paragraph()
    p.add_run().add_picture(path, width=Inches(width))
    cap = doc.add_paragraph()
    r = cap.add_run(caption)
    r.font.size = Pt(9)
    r.italic = True
    return p


def add_df_table(doc, df, caption, floatfmt='{:.2f}', caption_above=True):
    if caption_above:
        cap = doc.add_paragraph()
        r = cap.add_run(caption); r.bold = True; r.font.size = Pt(10)
    t = doc.add_table(rows=1 + len(df), cols=len(df.columns))
    t.style = 'Light Grid Accent 1'
    for j, c in enumerate(df.columns):
        t.rows[0].cells[j].text = str(c)
    for i, (_, row) in enumerate(df.iterrows()):
        for j, v in enumerate(row):
            t.rows[i + 1].cells[j].text = (floatfmt.format(v)
                                         if isinstance(v, float) else str(v))
    return t


def ref_runs(par, text_with_refs):
    """Add text where [n] citation markers become superscript runs."""
    for chunk in re.split(r'(\[\d+(?:[–,-]\d+)*\])', text_with_refs):
        if re.fullmatch(r'\[\d+(?:[–,-]\d+)*\]', chunk):
            r = par.add_run(chunk.strip('[]').replace('-', '–'))
            r.font.superscript = True
        else:
            par.add_run(chunk)
