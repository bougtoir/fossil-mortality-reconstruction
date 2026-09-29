"""SUPPLEMENT_Paleobiology.docx — supplementary material in journal style
(Supplementary Table/Figure numbering, TNR, supplementary files are not
typeset so layout is simple)."""
import pandas as pd, os, sys
sys.path.insert(0, '.')
from docx import Document
from docx.shared import Pt

TAB = '../output/tables'
OUT = '../manuscript'
rprof = pd.read_csv(f'{TAB}/identifiability_r_profile.csv')

doc = Document()
st = doc.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(12)
doc.add_paragraph('Supplementary Material').runs[0].bold = True
doc.add_paragraph('for: Beyond the oldest known individual: reconstructing '
                  'mortality distributions from fossil age-at-death evidence')

doc.add_heading('S1. Likelihood terms', level=1)
doc.add_paragraph('Individual ages (exact): L = Σ log g(a_i). Interval-'
 'censored: L = Σ log[G(hi_i) − G(lo_i)]. Lower bound: L = Σ log[1 − G(lo_i)].')
doc.add_paragraph('Grouped counts over bins [e_i, e_i+1): multinomial with '
 'bin masses G(e_i+1) − G(e_i).')
doc.add_paragraph('Maximum-only (m, n): L = log n + log g(m) + (n − 1) '
 'log G(m). Applied only when the maximum derives from a different sample '
 'than the individual ages; a maximum inside the modelled sample is used '
 'only as a posterior-predictive check.')

doc.add_heading('S2. Supplementary Table 1', level=1)
cp = doc.add_paragraph(); cp.add_run(
 'Supplementary Table 1. Profile of fitted Weibull mean/median age at '
 'death (yr) as the assumed standing-crop growth rate r varies; truth '
 'r = −0.05, n = 80.').bold = True
sub = rprof[['r', 'median', 'mean']].round(2)
t = doc.add_table(rows=1 + len(sub), cols=3); t.style = 'Table Grid'
for j, c in enumerate(sub.columns): t.rows[0].cells[j].text = c
for i, (_, row) in enumerate(sub.iterrows()):
    for j, v in enumerate(row): t.rows[i + 1].cells[j].text = str(v)

doc.add_heading('S3. Data statement', level=1)
doc.add_paragraph('Sources that could not be lawfully/reliably fetched '
 '(publisher TDM blocks) are marked unusable in data_inventory.csv; '
 'nothing was fabricated. Persisted raw inputs with URL, timestamp and '
 'SHA-256 are in data/raw/provenance.csv.')

doc.add_heading('S4. Model families and priors', level=1)
doc.add_paragraph('Weibull (λ, k); Gompertz–Makeham (a, b, c); Siler '
 '(a1, b1, a2, a3, b3); piecewise-constant hazard (5 bins). Fitting: '
 'Nelder–Mead MLE on log parameters, 4–6 jittered restarts, weakly '
 'informative prior sd = 2.5. Bayesian demo: emcee, 24 walkers, 5000 '
 'steps, marginal likelihoods unnormalized; posteriors summarized as '
 'median [95% CrI]. All random draws use crc32-based deterministic seeds '
 'per experiment cell.')

doc.save(f'{OUT}/SUPPLEMENT_Paleobiology.docx')
print('supplement saved')
