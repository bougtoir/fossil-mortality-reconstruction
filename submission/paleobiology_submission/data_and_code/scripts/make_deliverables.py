"""Generate remaining deliverables: figures.pptx, COVER_LETTER.docx,
SUPPLEMENT.docx."""
import pandas as pd, os, sys
sys.path.insert(0, '.')
from pptx import Presentation
from pptx.util import Inches, Pt
from docx import Document
from docx.shared import Pt as DPt, Inches as DInches
from docx_utils import add_equation, add_df_table

OUT = '../manuscript'; FIG = '../output/figures'; TAB = '../output/tables'
os.makedirs(OUT, exist_ok=True)

FIGS = [
    ('fig_scheme.png', 'Figure 1. Generative chain: latent mortality → '
     'assemblage mechanism → observation q(a) → age estimation → report type.'),
    ('fig_identifiability.png', 'Figure 2. Identifiability exhibits: q '
     'misspecification surface; r-dependence of catastrophic estimates; '
     'max-only likelihood ridge.'),
    ('fig_recoverability_map.png', 'Figure 3. Recoverability map: median '
     'age-at-death bias across fossilization scenarios and sample sizes.'),
    ('fig_recoverability_classes.png', 'Figure 4. Recoverability class '
     'distribution (A–D) by mortality summary quantity.'),
    ('fig_maximum.png', 'Figure 5. Evidence models for the maximum; E[M_n] '
     'under latent f vs observed g.'),
    ('fig_fossil_fits.png', 'Figure 6. Maiasaura demonstration: data, '
     'survival fits, sensitivity of the median estimate.'),
    ('fig_posterior.png', 'Figure 7. Bayesian posterior for Maiasaura '
     'latent mortality and standardized maximum.'),
]

prs = Presentation()
prs.slide_width = Inches(10); prs.slide_height = Inches(7.5)
for f, cap in FIGS:
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(Inches(0.3), Inches(0.2), Inches(9.4), Inches(0.6))
    tb.text_frame.text = cap
    s.shapes.add_picture(f'{FIG}/{f}', Inches(0.5), Inches(0.9),
                         height=Inches(5.8))
prs.save(f'{OUT}/figures.pptx')

# ---------------- COVER LETTER
doc = Document()
doc.add_paragraph('Dear Editors,')
for t in [
 'We submit "Beyond the oldest known individual: sample-size-aware '
 'reconstruction of mortality in extinct taxa" for consideration as an '
 'article.',
 'The paper asks a question the palaeobiological literature has lived '
 'with for a century: what does fossil age-at-death evidence actually '
 'license us to say about the mortality of extinct species? Our answer is '
 'a generative framework that separates the latent mortality hazard from '
 'assemblage formation, age-dependent preservation and recovery, and '
 'age-measurement error, together with a simulation-validated map of which '
 'mortality summaries survive fossilization.',
 'Three contributions are, to our knowledge, new: (i) an order-statistic '
 'treatment of the "oldest known specimen" that turns maxima into a '
 'likelihood term rather than a headline; (ii) a joint likelihood over '
 'individual ages, grouped tables, and max-only reports that avoids double '
 'counting; and (iii) a recoverability map built by degrading 45 extant '
 'mammal life tables into fossil-like samples. Demonstrations on '
 'Maiasaura, Psittacosaurus, Albertosaurus, and mammaliaform cementum '
 'series are reported as scenario ranges, and where the data cannot '
 'support an estimate we say so explicitly.',
 'All raw inputs, code, seeds, and generated outputs are included for '
 'reproduction; numbers in the manuscript are computed from the supplied '
 'result files, not transcribed.',
 'Yours faithfully,',
 'The authors']:
    doc.add_paragraph(t)
doc.save(f'{OUT}/COVER_LETTER.docx')

# ---------------- SUPPLEMENT
doc = Document()
doc.add_heading('Supplementary Information', 0)
doc.add_paragraph('"Beyond the oldest known individual" — methods detail, '
                  'additional tables, and data statement.')
doc.add_heading('S1. Likelihood terms', 1)
doc.add_paragraph('Individual ages (exact): L = Σ log g(a_i). Interval-censored: '
                  'L = Σ log[G(hi_i) − G(lo_i)]. Lower bounds: L = Σ log[1 − G(lo_i)].')
doc.add_paragraph('Grouped counts over bins [e_i, e_i+1): multinomial with '
                  'bin masses G(e_i+1) − G(e_i).')
doc.add_paragraph('Maximum-only (m, n): L = log n + log g(m) + (n − 1) log G(m). '
                  'Applied only when the maximum derives from an independent '
                  'specimen pool; a maximum of the modelled dataset itself '
                  'enters only via posterior-predictive checks.')
doc.add_heading('S2. Identifiability tables', 1)
rprof = pd.read_csv(f'{TAB}/identifiability_r_profile.csv')
doc.add_paragraph('Table S1. Profile of fitted Weibull median/mean age at '
                  'death over assumed growth rate r (catastrophic scenario, '
                  'n = 80; true r = −0.05).')
add_df_table(doc, rprof.round(3), 'Table S1 (data).', '{:.3f}', caption_above=False)
doc.add_heading('S3. Data statement', 1)
inv = pd.read_csv('../data_inventory.csv')
tab = inv[['taxon', 'source_citation', 'fossil_or_extant', 'n_aged',
           'age_data_type', 'usable_primary', 'exclusion_reason']]
add_df_table(doc, tab, 'Table S2. Dataset inventory (abbreviated; full '
                       'machine-readable version in data_inventory.csv).',
             '{:.0f}', caption_above=False)
doc.add_paragraph('Sources that could not be lawfully/reliably fetched '
                  '(publisher TDM blocks) are marked unusable; nothing was '
                  'fabricated in their place. Raw files and a provenance '
                  'ledger (URL, UTC timestamp, SHA-256, license) are in '
                  'data/raw/.')
doc.add_heading('S4. Model families and priors', 1)
doc.add_paragraph('Weibull (λ, k); Gompertz–Makeham (a, b, c); Siler '
                  '(a1, b1, a2, a3, b3); piecewise-constant hazard (5 bins). '
                  'All parameters positive, fitted on the log scale. '
                  'Weakly-informative log-normal priors (sd 2.5) used in '
                  'MAP/penalty fits; prior sensitivity reported in '
                  'phase7_sensitivity.csv.')
doc.save(f'{OUT}/SUPPLEMENT.docx')
print('done')
